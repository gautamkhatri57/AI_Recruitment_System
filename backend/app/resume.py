import io
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document


def extract_text_from_pdf(file_bytes):
    text = ""

    try:
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

    except Exception as e:
        raise ValueError(
            f"PDF extraction failed: {str(e)}"
        )

    return text.strip()


def extract_text_from_docx(file_bytes):
    text = ""

    try:
        docx_file = io.BytesIO(file_bytes)
        document = Document(docx_file)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"

    except Exception as e:
        raise ValueError(
            f"DOCX extraction failed: {str(e)}"
        )

    return text.strip()


def extract_text_from_file(file_bytes, filename):
    if not file_bytes:
        raise ValueError("Uploaded file is empty.")

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)

    if extension == ".docx":
        return extract_text_from_docx(file_bytes)

    raise ValueError(
        "Unsupported file format. Please upload PDF or DOCX."
    )